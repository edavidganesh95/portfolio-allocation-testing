"""Table presentation layer.

Two jobs:

1. **One formatting contract.** A column called ``CAGR`` renders as ``12.3%``
   everywhere it appears — in the app, in the Excel export and in the PDF —
   because the format is resolved from the column name, not restated at each
   call site.
2. **Readable density.** Institutional tables carry a lot of numbers, so
   magnitude is encoded as a quiet background wash and sign as colour. The
   reader can find the big and the negative numbers without reading every cell.

Values are never rounded before display: formatting is applied by the Styler,
so exports and charts continue to see full precision.
"""

from __future__ import annotations

import io
import re
from collections.abc import Callable, Iterable

import numpy as np
import pandas as pd

from . import theme

PCT0 = "{:.0%}"
PCT1 = "{:.1%}"
PCT2 = "{:.2%}"
SIGNED_PCT1 = "{:+.1%}"
RATIO2 = "{:.2f}"
MONEY0 = "${:,.0f}"
COUNT = "{:,.0f}"
SCORE0 = "{:.0f}"

#: Column-name patterns to display formats, most specific first.
_FORMAT_RULES: tuple[tuple[str, str], ...] = (
    (r"^Trade$", SIGNED_PCT1),
    (r"Wealth$", MONEY0),
    (
        r"^Runs$|^Rebalances$|^Observations$|^Period$|Months$|^Holdings Above|^Actual ETFs$|^Blocks$",
        COUNT,
    ),
    (r"^Stability Score$", SCORE0),
    (
        r"^Sharpe$|^Expected Sharpe$|^Sortino$|^HHI$|^Effective Holdings$|^Effective ETFs$|^Effective Risk Bets$|^Diversification Ratio$|Correlation",
        RATIO2,
    ),
    (r"^Inclusion Frequency$|Win Rate$|^Chance ", PCT0),
    (
        r"CAGR|E\(r\)|Volatility$|Drawdown$|^CVaR|Weight|Percentile|^Minimum$|^Maximum$"
        r"|Turnover|Reallocate|^P\(|Target$|^Robust |Share$|Return$|Allocation$|^Largest Holding$",
        PCT1,
    ),
)

#: Columns where a higher number is worse, so the wash must run the other way.
_INVERSE_COLUMNS = (
    "Volatility",
    "Historical Volatility",
    "Expected Volatility",
    "Max Drawdown",
    "Weight Std Dev",
    "Weighted Correlation",
    "Largest Risk Share",
    "Largest Cluster Allocation",
    "Capital to Reallocate",
)


def display_format(column: str) -> str | None:
    """Resolve the display format for a column name."""
    for pattern, spec in _FORMAT_RULES:
        if re.search(pattern, str(column)):
            return spec
    return None


def formats_for(
    frame: pd.DataFrame,
    default: str | None = None,
) -> dict[str, str]:
    """Format map covering every numeric column of ``frame``.

    ``default`` covers frames whose columns are data rather than metric names —
    a method x ticker weight matrix, or a correlation matrix — where the column
    label cannot imply a format.
    """
    out: dict[str, str] = {}
    for column in frame.columns:
        if not pd.api.types.is_numeric_dtype(frame[column]):
            continue
        spec = display_format(column) or default
        if spec:
            out[column] = spec
    return out


# ----------------------------------------------------------------------
# Cell styling
# ----------------------------------------------------------------------


def _column_wash(
    series: pd.Series,
    palette: tuple[str, ...],
    invert: bool,
    strength: float,
) -> list[str]:
    values = pd.to_numeric(series, errors="coerce")
    finite = values.replace([np.inf, -np.inf], np.nan).dropna()
    if finite.empty:
        return [""] * len(series)

    low = float(finite.min())
    high = float(finite.max())
    if not np.isfinite(low) or not np.isfinite(high) or high - low <= 1e-15:
        return [""] * len(series)

    styles = []
    for value in values:
        if pd.isna(value) or not np.isfinite(value):
            styles.append("")
            continue
        position = (float(value) - low) / (high - low)
        if invert:
            position = 1.0 - position
        color = theme.mix(
            theme.SURFACE,
            theme.ramp(palette, position),
            strength * position,
        )
        # The wash can run dark enough to swallow ink, so the text colour is
        # chosen per cell rather than left to the default.
        text = theme.readable_text(color)
        styles.append(f"background-color: {color}; color: {text}")
    return styles


def heat(
    styler: pd.io.formats.style.Styler,
    columns: Iterable[str],
    palette: tuple[str, ...] = theme.SEQUENTIAL,
    strength: float = 0.62,
) -> pd.io.formats.style.Styler:
    """Quiet magnitude wash behind the named columns."""
    frame = styler.data
    present = [c for c in columns if c in frame.columns]
    for column in present:
        invert = any(key in column for key in _INVERSE_COLUMNS)
        styler = styler.apply(
            _column_wash,
            palette=palette,
            invert=invert,
            strength=strength,
            subset=[column],
        )
    return styler


def _sign_style(series: pd.Series) -> list[str]:
    values = pd.to_numeric(series, errors="coerce")
    out = []
    for value in values:
        if pd.isna(value):
            out.append("")
        elif value > 1e-12:
            out.append(f"color: {theme.POSITIVE}; font-weight: 600")
        elif value < -1e-12:
            out.append(f"color: {theme.NEGATIVE}; font-weight: 600")
        else:
            out.append(f"color: {theme.MUTED}")
    return out


def sign(
    styler: pd.io.formats.style.Styler,
    columns: Iterable[str],
) -> pd.io.formats.style.Styler:
    """Colour positive values green and negative values brick."""
    present = [c for c in columns if c in styler.data.columns]
    if present:
        styler = styler.apply(_sign_style, subset=present)
    return styler


_VERDICT_STYLES = {
    "PASS": f"color: {theme.POSITIVE}; font-weight: 700",
    "FAIL": f"color: {theme.NEGATIVE}; font-weight: 700",
    "N/A": f"color: {theme.MUTED}",
    "HOLD": f"color: {theme.MUTED}; font-weight: 600",
    "INCREASE": f"color: {theme.POSITIVE}; font-weight: 700",
    "REDUCE": f"color: {theme.NEGATIVE}; font-weight: 700",
    "True": f"color: {theme.POSITIVE}; font-weight: 600",
    "False": f"color: {theme.NEGATIVE}; font-weight: 600",
}


def verdict(
    styler: pd.io.formats.style.Styler,
    columns: Iterable[str],
) -> pd.io.formats.style.Styler:
    """Colour categorical verdict columns (PASS/FAIL, INCREASE/REDUCE)."""
    present = [c for c in columns if c in styler.data.columns]
    for column in present:
        styler = styler.map(
            lambda value: _VERDICT_STYLES.get(str(value), ""),
            subset=[column],
        )
    return styler


def highlight_rows(
    styler: pd.io.formats.style.Styler,
    labels: Iterable[str],
    color: str = theme.BRASS_SOFT,
) -> pd.io.formats.style.Styler:
    """Tint whole rows — used to keep the reference portfolio findable."""
    wanted = {str(label) for label in labels}

    def _row(row: pd.Series) -> list[str]:
        if str(row.name) in wanted:
            return [f"background-color: {color}"] * len(row)
        return [""] * len(row)

    return styler.apply(_row, axis=1)


def style(
    frame: pd.DataFrame,
    heat_columns: Iterable[str] | None = None,
    sign_columns: Iterable[str] | None = None,
    verdict_columns: Iterable[str] | None = None,
    highlight: Iterable[str] | None = None,
    palette: tuple[str, ...] = theme.SEQUENTIAL,
    formats: dict[str, str] | None = None,
    default_format: str | None = None,
    precision_overrides: dict[str, str] | None = None,
) -> pd.io.formats.style.Styler:
    """Build the standard research-table Styler for ``frame``."""
    resolved = (
        dict(formats)
        if formats is not None
        else formats_for(frame, default=default_format)
    )
    if precision_overrides:
        resolved.update(precision_overrides)

    styler = frame.style.format(resolved, na_rep="—")

    if highlight:
        styler = highlight_rows(styler, highlight)
    if heat_columns:
        styler = heat(styler, heat_columns, palette=palette)
    if sign_columns:
        styler = sign(styler, sign_columns)
    if verdict_columns:
        styler = verdict(styler, verdict_columns)
    return styler


def metric_columns(frame: pd.DataFrame) -> list[str]:
    """Numeric columns worth washing in a return/risk summary table."""
    candidates = (
        "CAGR",
        "Volatility",
        "Sharpe",
        "Sortino",
        "Max Drawdown",
    )
    columns = [c for c in frame.columns if c in candidates]
    columns += [c for c in frame.columns if str(c).startswith("CVaR")]
    return columns


# ----------------------------------------------------------------------
# Streamlit rendering
# ----------------------------------------------------------------------


def render(
    data,
    height: int | None = None,
    hide_index: bool = False,
    column_config: dict | None = None,
) -> None:
    """Render a table at full container width with current Streamlit kwargs."""
    import streamlit as st

    kwargs: dict = {"width": "stretch"}
    if height is not None:
        kwargs["height"] = height
    if hide_index:
        kwargs["hide_index"] = True
    if column_config:
        kwargs["column_config"] = column_config
    st.dataframe(data, **kwargs)


def weight_column_config(
    columns: Iterable[str],
    max_value: float = 1.0,
) -> dict:
    """Progress-bar columns for allocation weights."""
    import streamlit as st

    return {
        column: st.column_config.ProgressColumn(
            column,
            format="percent",
            min_value=0.0,
            max_value=float(max_value),
        )
        for column in columns
    }


# ----------------------------------------------------------------------
# Exports
# ----------------------------------------------------------------------


def to_csv_bytes(frame: pd.DataFrame) -> bytes:
    return frame.to_csv().encode("utf-8-sig")


def _sheet_name(name: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", " ", str(name)).strip()[:31] or "Sheet"
    candidate = cleaned
    suffix = 2
    while candidate.lower() in used:
        tail = f" {suffix}"
        candidate = cleaned[: 31 - len(tail)] + tail
        suffix += 1
    used.add(candidate.lower())
    return candidate


def to_excel_bytes(
    tables: dict[str, pd.DataFrame],
    number_formats: Callable[[str], str | None] = display_format,
) -> bytes:
    """Workbook with one formatted sheet per table.

    Excel receives the underlying numbers with a matching number format, so
    the recipient can keep working with the data instead of re-typing strings.
    """
    excel_formats = {
        PCT0: "0%",
        PCT1: "0.0%",
        PCT2: "0.00%",
        SIGNED_PCT1: "+0.0%;-0.0%",
        RATIO2: "0.00",
        MONEY0: "$#,##0",
        COUNT: "#,##0",
        SCORE0: "0",
    }

    from openpyxl.styles import Font, PatternFill

    buffer = io.BytesIO()
    used: set[str] = set()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, frame in tables.items():
            if frame is None or not isinstance(frame, pd.DataFrame) or frame.empty:
                continue
            sheet = _sheet_name(name, used)
            frame.to_excel(writer, sheet_name=sheet)
            worksheet = writer.sheets[sheet]

            index_width = max(
                [len(str(frame.index.name or ""))]
                + [len(str(v)) for v in frame.index[:400]]
                + [10]
            )
            worksheet.column_dimensions["A"].width = min(index_width + 3, 42)
            worksheet.freeze_panes = "B2"

            for offset, column in enumerate(frame.columns, start=2):
                letter = worksheet.cell(row=1, column=offset).column_letter
                worksheet.column_dimensions[letter].width = min(
                    max(len(str(column)) + 3, 11), 30
                )
                spec = number_formats(column)
                fmt = excel_formats.get(spec) if spec else None
                if not fmt:
                    continue
                for row in range(2, len(frame) + 2):
                    worksheet.cell(row=row, column=offset).number_format = fmt

            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill("solid", fgColor=theme.NAVY.lstrip("#"))
            for cell in worksheet[1]:
                cell.font = header_font
                cell.fill = header_fill

    return buffer.getvalue()
