import numpy as np
import pandas as pd
import pytest

from src import tables, theme

#: Every column name the six sections put on screen, with the format the
#: research output is specified to use. This is the contract the app relies on
#: when it stops passing explicit format dictionaries.
EXPECTED_FORMATS = {
    "CAGR": tables.PCT1,
    "Volatility": tables.PCT1,
    "Sharpe": tables.RATIO2,
    "Sortino": tables.RATIO2,
    "Max Drawdown": tables.PCT1,
    "CVaR 90%": tables.PCT1,
    "CVaR 95%": tables.PCT1,
    "CVaR 99%": tables.PCT1,
    "Median Weight": tables.PCT1,
    "25th Percentile": tables.PCT1,
    "75th Percentile": tables.PCT1,
    "Minimum": tables.PCT1,
    "Maximum": tables.PCT1,
    "Inclusion Frequency": tables.PCT0,
    "Weight Std Dev": tables.PCT1,
    "Stability Score": tables.SCORE0,
    "Runs": tables.COUNT,
    "Consensus Weight": tables.PCT1,
    "Median Terminal Wealth": tables.MONEY0,
    "5th Percentile Wealth": tables.MONEY0,
    "25th Percentile Wealth": tables.MONEY0,
    "75th Percentile Wealth": tables.MONEY0,
    "95th Percentile Wealth": tables.MONEY0,
    "Median CAGR": tables.PCT1,
    "5th Percentile CAGR": tables.PCT1,
    "P(Loss at Horizon)": tables.PCT1,
    "Median Max Drawdown": tables.PCT1,
    "5th Percentile Max Drawdown": tables.PCT1,
    "P(CAGR > 6%)": tables.PCT1,
    "P(CAGR > 8%)": tables.PCT1,
    "P(CAGR > 10%)": tables.PCT1,
    "Block Win Rate": tables.PCT0,
    "Average Rebalance Turnover": tables.PCT1,
    "Median Rebalance Turnover": tables.PCT1,
    "Maximum Rebalance Turnover": tables.PCT1,
    "Rebalances": tables.COUNT,
    "Current Weight": tables.PCT1,
    "Robust Lower": tables.PCT1,
    "Robust Median": tables.PCT1,
    "Robust Upper": tables.PCT1,
    "Minimal Target": tables.PCT1,
    "Trade": tables.SIGNED_PCT1,
    "Expected Return": tables.PCT1,
    "Correlation": tables.RATIO2,
    "Average Correlation": tables.RATIO2,
    "Risk Share": tables.PCT1,
    "Observations": tables.COUNT,
    "Train Months": tables.COUNT,
    "Test Months": tables.COUNT,
}


@pytest.mark.parametrize("column,expected", sorted(EXPECTED_FORMATS.items()))
def test_display_format_contract(column, expected):
    assert tables.display_format(column) == expected


def test_unknown_columns_have_no_forced_format():
    assert tables.display_format("Message") is None
    assert tables.display_format("Ticker") is None


def test_formats_for_skips_non_numeric_columns():
    frame = pd.DataFrame({"CAGR": [0.1], "Message": ["ok"]})
    assert tables.formats_for(frame) == {"CAGR": tables.PCT1}


@pytest.fixture
def metrics_frame():
    return pd.DataFrame(
        {
            "CAGR": [0.201, 0.09, 0.137],
            "Volatility": [0.192, 0.162, 0.158],
            "Sharpe": [1.05, 0.55, 0.87],
            "Max Drawdown": [-0.309, -0.319, -0.257],
            "CVaR 95%": [0.094, 0.093, 0.061],
        },
        index=pd.Index(["TMFC", "VWO", "ACWI"], name="Ticker"),
    )


def test_style_renders_formatted_values(metrics_frame):
    html = tables.style(metrics_frame).to_html()
    assert "20.1%" in html
    assert "1.05" in html


def test_style_applies_a_background_wash(metrics_frame):
    html = tables.style(
        metrics_frame, heat_columns=tables.metric_columns(metrics_frame)
    ).to_html()
    assert "background-color" in html


def test_heat_wash_inverts_for_risk_columns(metrics_frame):
    """A high volatility must not be washed like a high return."""
    styler = tables.style(metrics_frame, heat_columns=["CAGR", "Volatility"])
    rendered = styler._compute().ctx

    positions = {label: i for i, label in enumerate(metrics_frame.index)}
    cagr_col = list(metrics_frame.columns).index("CAGR")
    vol_col = list(metrics_frame.columns).index("Volatility")

    def background(row: int, col: int) -> str:
        return dict(rendered[(row, col)])["background-color"]

    # TMFC has both the highest CAGR and the highest volatility. The CAGR cell
    # should be the darkest in its column and the volatility cell the lightest,
    # because a large drawdown or volatility is not an achievement.
    assert background(positions["TMFC"], cagr_col) != theme.SURFACE
    assert background(positions["TMFC"], vol_col) == theme.SURFACE
    assert background(positions["ACWI"], vol_col) != theme.SURFACE


def test_metric_columns_picks_up_any_cvar_level(metrics_frame):
    assert "CVaR 95%" in tables.metric_columns(metrics_frame)


def test_sign_colours_positive_and_negative():
    frame = pd.DataFrame({"Trade": [0.05, -0.05, 0.0]}, index=["A", "B", "C"])
    html = tables.style(frame, sign_columns=["Trade"]).to_html()
    assert theme.POSITIVE in html
    assert theme.NEGATIVE in html
    assert "+5.0%" in html and "-5.0%" in html


def test_verdict_colours_pass_and_fail():
    frame = pd.DataFrame({"Result": ["PASS", "FAIL", "N/A"]}, index=["A", "B", "C"])
    html = tables.style(frame, verdict_columns=["Result"]).to_html()
    assert theme.POSITIVE in html
    assert theme.NEGATIVE in html


def test_highlight_tints_only_the_named_rows(metrics_frame):
    styler = tables.style(metrics_frame, highlight=["VWO"])
    rendered = styler._compute().ctx
    tinted = {row for (row, _col), styles in rendered.items() if styles}
    assert tinted == {list(metrics_frame.index).index("VWO")}


def test_style_renders_missing_values_as_an_em_dash():
    frame = pd.DataFrame({"CAGR": [np.nan]}, index=["A"])
    assert "—" in tables.style(frame).to_html()


def test_heat_ignores_constant_columns():
    frame = pd.DataFrame({"CAGR": [0.1, 0.1]}, index=["A", "B"])
    rendered = tables.style(frame, heat_columns=["CAGR"])._compute().ctx
    assert not any(rendered.values())


def test_csv_export_round_trips(metrics_frame):
    payload = tables.to_csv_bytes(metrics_frame)
    restored = pd.read_csv(pd.io.common.BytesIO(payload), index_col=0)
    pd.testing.assert_frame_equal(restored, metrics_frame, check_names=False)


def test_excel_export_keeps_numbers_and_number_formats(metrics_frame):
    import openpyxl

    payload = tables.to_excel_bytes({"Diagnostics": metrics_frame})
    book = openpyxl.load_workbook(pd.io.common.BytesIO(payload))
    sheet = book["Diagnostics"]

    assert sheet["A1"].value == "Ticker"
    # Values stay numeric so the recipient can keep working with them.
    assert sheet["B2"].value == pytest.approx(0.201)
    assert sheet["B2"].number_format == "0.0%"
    assert sheet["D2"].number_format == "0.00"
    assert sheet.freeze_panes == "B2"


def test_excel_export_skips_empty_frames():
    import openpyxl

    payload = tables.to_excel_bytes(
        {"Empty": pd.DataFrame(), "Real": pd.DataFrame({"CAGR": [0.1]})}
    )
    book = openpyxl.load_workbook(pd.io.common.BytesIO(payload))
    assert book.sheetnames == ["Real"]


def test_excel_export_deduplicates_and_sanitises_sheet_names():
    import openpyxl

    frame = pd.DataFrame({"CAGR": [0.1]})
    payload = tables.to_excel_bytes(
        {
            "Return / risk [2024]": frame,
            "A very long section name that exceeds the Excel limit": frame,
            "A very long section name that exceeds the Excel limit ": frame,
        }
    )
    book = openpyxl.load_workbook(pd.io.common.BytesIO(payload))
    assert len(book.sheetnames) == 3
    assert all(len(name) <= 31 for name in book.sheetnames)
    assert not any(set(name) & set("[]:*?/\\") for name in book.sheetnames)


def test_default_format_covers_data_named_columns():
    """A method x ticker weight matrix has no metric names to infer from."""
    weights = pd.DataFrame(
        [[0.090909, 0.5]], index=["Equal Weight"], columns=["TMFC", "VWO"]
    )
    assert tables.formats_for(weights) == {}
    assert tables.formats_for(weights, default=tables.PCT1) == {
        "TMFC": tables.PCT1,
        "VWO": tables.PCT1,
    }

    html = tables.style(weights, default_format=tables.PCT1).to_html()
    assert "9.1%" in html
    assert "0.090909" not in html


def test_named_columns_still_win_over_the_default():
    frame = pd.DataFrame({"Sharpe": [1.05], "TMFC": [0.25]})
    resolved = tables.formats_for(frame, default=tables.PCT1)
    assert resolved == {"Sharpe": tables.RATIO2, "TMFC": tables.PCT1}
